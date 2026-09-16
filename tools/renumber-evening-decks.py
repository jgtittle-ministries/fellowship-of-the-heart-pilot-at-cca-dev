# -*- coding: utf-8 -*-
"""Master renumbering: one continuous sequence across all 22 evening decks.

Page 1 = the shared series title slide (every deck's slide 1).
Every subsequent slide, session by session in order, takes the next integer.
Letter-suffix interim numbers are retired.
"""
import sys

sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.oxml.ns import qn

BASE = r"_implementation-notes/evening-series"

counter = 2
ranges = {}
for n in range(1, 23):
    path = BASE + r"\FotH Evening Session %d Slides DRAFT v1.pptx" % n
    p = Presentation(path)
    start = counter
    for i, s in enumerate(p.slides):
        target = "1" if i == 0 else str(counter)
        numbox = None
        for sh in s.shapes:
            if sh.name == "SlideNum":
                numbox = sh
                break
        if numbox is None:
            import re
            for sh in s.shapes:
                if (sh.has_text_frame and sh.left is not None and sh.left > 10500000
                        and re.fullmatch(r"\d+a?", sh.text_frame.text.strip())):
                    numbox = sh
                    break
        if numbox is None:
            raise SystemExit("S%d slide %d has no SlideNum box" % (n, i + 1))
        ts = list(numbox._element.iter(qn("a:t")))
        ts[0].text = target
        for t in ts[1:]:
            t.text = ""
        if i > 0:
            counter += 1
    p.save(path)
    ranges[n] = (start, counter - 1)
    print("S%-2d -> pages %d–%d" % (n, start, counter - 1))

print("last page:", counter - 1)

# numbering map for the repo
lines = [
    "# Evening Deck Master Numbering — of record",
    "",
    "*Rebuilt 2026-09-16 at John's word: one continuous sequence across all",
    "22 session decks, replacing the pre-ripple 1–277 numbering and its",
    "interim letter-suffix pages. Page 1 is the shared series title slide",
    "(every deck's first slide). Corner numbers are per-slide text boxes;",
    "renumber by rerunning the renumber script after any future insert.*",
    "",
    "| Session | Pages |",
    "|---|---|",
]
for n in range(1, 23):
    a, b = ranges[n]
    lines.append("| %d | %d–%d |" % (n, a, b))
lines.append("")
lines.append("Total: %d master pages (1 + %d session pages)." % (counter - 1, counter - 2))
lines.append("")
with open(BASE + r"\slide-numbering-map.md", "w", encoding="utf-8", newline="\n") as f:
    f.write("\n".join(lines))
print("map written")
